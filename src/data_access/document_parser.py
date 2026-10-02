from __future__ import annotations

import json
import re
from pathlib import Path
from typing import Any, Dict, List, Optional

from pypdf import PdfReader
from pptx import Presentation
from docx import Document

try:
    import extract_msg
except ImportError:
    extract_msg = None


# ============================================================
# PROJECT PATHS
# ============================================================

PROJECT_ROOT = Path(__file__).resolve().parents[2]
DATA_DIR = PROJECT_ROOT / "data"
DOCUMENT_DIR = DATA_DIR / "documents"


SUPPORTED_EXTENSIONS = {
    ".pdf",
    ".pptx",
    ".docx",
    ".msg",
}


# ============================================================
# COMMON HELPERS
# ============================================================

def _relative_source(path: Path) -> str:
    """
    Return project-relative path for provenance.
    """
    try:
        return str(path.resolve().relative_to(PROJECT_ROOT.resolve()))
    except ValueError:
        return str(path)


def _clean_text(text: Any) -> str:
    """
    Normalize whitespace without destroying paragraph boundaries.
    """
    if text is None:
        return ""

    text = str(text)
    text = text.replace("\x00", " ")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)

    return text.strip()


def _citation(source: str, locator: str) -> str:
    """
    Canonical evidence citation.
    """
    return f"{source}#{locator}"


def _base_evidence(
    source_path: Path,
    document_type: str,
    locator: str,
    content_type: str,
    text: str = "",
) -> Dict[str, Any]:

    source = _relative_source(source_path)

    return {
        "source": source,
        "document_type": document_type,
        "locator": locator,
        "content_type": content_type,
        "text": _clean_text(text),
        "citation": _citation(source, locator),
    }


# ============================================================
# PDF
# ============================================================

def parse_pdf(source_path: str | Path) -> Dict[str, Any]:
    """
    Parse PDF into page-level evidence.

    Captures:
      - page text
      - page number
      - image count
      - basic page metadata

    Tables/charts that are embedded as images are currently
    represented as image objects rather than OCR'd.
    """

    path = Path(source_path)
    reader = PdfReader(str(path))

    evidence = []

    for page_number, page in enumerate(reader.pages, start=1):

        try:
            text = page.extract_text() or ""
        except Exception:
            text = ""

        image_count = 0

        try:
            image_count = len(page.images)
        except Exception:
            image_count = 0

        page_evidence = _base_evidence(
            source_path=path,
            document_type="pdf",
            locator=f"page={page_number}",
            content_type="page",
            text=text,
        )

        page_evidence["page_number"] = page_number
        page_evidence["image_count"] = image_count

        evidence.append(page_evidence)

    return {
        "document": _relative_source(path),
        "document_type": "pdf",
        "page_count": len(reader.pages),
        "evidence": evidence,
    }


# ============================================================
# PPTX
# ============================================================

def _extract_pptx_table(table) -> str:
    """
    Convert PowerPoint table into deterministic text.
    """

    rows = []

    for row in table.rows:

        cells = []

        for cell in row.cells:
            cells.append(_clean_text(cell.text))

        rows.append(" | ".join(cells))

    return "\n".join(rows)


def parse_pptx(source_path: str | Path) -> Dict[str, Any]:
    """
    Parse PowerPoint into slide-level evidence.

    Captures:
      - text boxes
      - tables
      - chart/image placeholders
      - slide number
    """

    path = Path(source_path)
    presentation = Presentation(str(path))

    evidence = []

    for slide_number, slide in enumerate(
        presentation.slides,
        start=1,
    ):

        slide_text_blocks = []
        shape_items = []

        for shape_number, shape in enumerate(
            slide.shapes,
            start=1,
        ):

            # ------------------------------------------------
            # Text
            # ------------------------------------------------

            if hasattr(shape, "text"):

                text = _clean_text(shape.text)

                if text:

                    item = _base_evidence(
                        source_path=path,
                        document_type="pptx",
                        locator=f"slide={slide_number}&shape={shape_number}",
                        content_type="text",
                        text=text,
                    )

                    item["slide_number"] = slide_number
                    item["shape_number"] = shape_number

                    evidence.append(item)

                    slide_text_blocks.append(text)
                    shape_items.append(item)

            # ------------------------------------------------
            # Table
            # ------------------------------------------------

            if getattr(shape, "has_table", False):

                table_text = _extract_pptx_table(
                    shape.table
                )

                item = _base_evidence(
                    source_path=path,
                    document_type="pptx",
                    locator=(
                        f"slide={slide_number}"
                        f"&shape={shape_number}"
                        f"&type=table"
                    ),
                    content_type="table",
                    text=table_text,
                )

                item["slide_number"] = slide_number
                item["shape_number"] = shape_number

                evidence.append(item)

                shape_items.append(item)

            # ------------------------------------------------
            # Picture
            # ------------------------------------------------

            if shape.shape_type == 13:
                item = _base_evidence(
                    source_path=path,
                    document_type="pptx",
                    locator=(
                        f"slide={slide_number}"
                        f"&shape={shape_number}"
                        f"&type=image"
                    ),
                    content_type="image",
                )

                item["slide_number"] = slide_number
                item["shape_number"] = shape_number

                evidence.append(item)

                shape_items.append(item)

            # ------------------------------------------------
            # Chart
            # ------------------------------------------------

            if getattr(shape, "has_chart", False):

                chart_type = str(
                    getattr(
                        shape.chart,
                        "chart_type",
                        "unknown",
                    )
                )

                item = _base_evidence(
                    source_path=path,
                    document_type="pptx",
                    locator=(
                        f"slide={slide_number}"
                        f"&shape={shape_number}"
                        f"&type=chart"
                    ),
                    content_type="chart",
                    text=f"PowerPoint chart type: {chart_type}",
                )

                item["slide_number"] = slide_number
                item["shape_number"] = shape_number
                item["chart_type"] = chart_type

                evidence.append(item)

                shape_items.append(item)

        # ----------------------------------------------------
        # Slide-level summary
        # ----------------------------------------------------

        slide_summary = _base_evidence(
            source_path=path,
            document_type="pptx",
            locator=f"slide={slide_number}",
            content_type="slide",
            text="\n".join(slide_text_blocks),
        )

        slide_summary["slide_number"] = slide_number
        slide_summary["shape_count"] = len(slide.shapes)

        evidence.append(slide_summary)

    return {
        "document": _relative_source(path),
        "document_type": "pptx",
        "slide_count": len(presentation.slides),
        "evidence": evidence,
    }


# ============================================================
# DOCX
# ============================================================

def _extract_docx_table(table) -> str:
    """
    Convert Word table into deterministic text.
    """

    rows = []

    for row in table.rows:

        cells = []

        for cell in row.cells:
            cells.append(_clean_text(cell.text))

        rows.append(" | ".join(cells))

    return "\n".join(rows)


def parse_docx(source_path: str | Path) -> Dict[str, Any]:
    """
    Parse DOCX into paragraph/table/image-level evidence.
    """

    path = Path(source_path)
    document = Document(str(path))

    evidence = []

    # --------------------------------------------------------
    # Paragraphs
    # --------------------------------------------------------

    for paragraph_number, paragraph in enumerate(
        document.paragraphs,
        start=1,
    ):

        text = _clean_text(paragraph.text)

        if not text:
            continue

        style_name = ""

        try:
            style_name = paragraph.style.name
        except Exception:
            pass

        item = _base_evidence(
            source_path=path,
            document_type="docx",
            locator=f"paragraph={paragraph_number}",
            content_type="paragraph",
            text=text,
        )

        item["paragraph_number"] = paragraph_number
        item["style"] = style_name

        evidence.append(item)

    # --------------------------------------------------------
    # Tables
    # --------------------------------------------------------

    for table_number, table in enumerate(
        document.tables,
        start=1,
    ):

        table_text = _extract_docx_table(table)

        item = _base_evidence(
            source_path=path,
            document_type="docx",
            locator=f"table={table_number}",
            content_type="table",
            text=table_text,
        )

        item["table_number"] = table_number
        item["row_count"] = len(table.rows)

        evidence.append(item)

    # --------------------------------------------------------
    # Inline images
    # --------------------------------------------------------

    image_count = 0

    for paragraph in document.paragraphs:

        for run in paragraph.runs:

            drawing_elements = run._element.xpath(
                ".//a:blip"
            )

            image_count += len(drawing_elements)

    if image_count:

        item = _base_evidence(
            source_path=path,
            document_type="docx",
            locator="images",
            content_type="image",
            text=f"DOCX contains {image_count} embedded image(s).",
        )

        item["image_count"] = image_count

        evidence.append(item)

    return {
        "document": _relative_source(path),
        "document_type": "docx",
        "paragraph_count": len(document.paragraphs),
        "table_count": len(document.tables),
        "image_count": image_count,
        "evidence": evidence,
    }


# ============================================================
# MSG / OUTLOOK EMAIL
# ============================================================

def parse_msg(source_path: str | Path) -> Dict[str, Any]:
    """
    Parse Outlook MSG using extract-msg.

    Captures:
      - subject
      - sender
      - recipients
      - date
      - body
      - attachments

    msgforge is intentionally NOT used here because it creates
    MSG files but does not parse/read them.
    """

    if extract_msg is None:

        raise RuntimeError(
            "extract-msg is not installed. "
            "Run: pip install extract-msg"
        )

    path = Path(source_path)

    message = extract_msg.Message(str(path))

    evidence = []

    # --------------------------------------------------------
    # Metadata
    # --------------------------------------------------------

    metadata = {
        "subject": getattr(message, "subject", None),
        "sender": getattr(message, "sender", None),
        "to": getattr(message, "to", None),
        "cc": getattr(message, "cc", None),
        "date": str(getattr(message, "date", None)),
    }

    metadata_text = (
        f"Subject: {metadata['subject']}\n"
        f"From: {metadata['sender']}\n"
        f"To: {metadata['to']}\n"
        f"CC: {metadata['cc']}\n"
        f"Date: {metadata['date']}"
    )

    metadata_item = _base_evidence(
        source_path=path,
        document_type="msg",
        locator="metadata",
        content_type="email_metadata",
        text=metadata_text,
    )

    metadata_item["metadata"] = metadata

    evidence.append(metadata_item)

    # --------------------------------------------------------
    # Body
    # --------------------------------------------------------

    body = getattr(message, "body", "") or ""

    body_item = _base_evidence(
        source_path=path,
        document_type="msg",
        locator="body",
        content_type="email_body",
        text=body,
    )

    evidence.append(body_item)

    # --------------------------------------------------------
    # Attachments
    # --------------------------------------------------------

    attachments = []

    for attachment_number, attachment in enumerate(
        getattr(message, "attachments", []) or [],
        start=1,
    ):

        filename = getattr(
            attachment,
            "longFilename",
            None,
        ) or getattr(
            attachment,
            "shortFilename",
            None,
        )

        attachment_info = {
            "attachment_number": attachment_number,
            "filename": filename,
        }

        attachments.append(attachment_info)

        attachment_item = _base_evidence(
            source_path=path,
            document_type="msg",
            locator=(
                f"attachment={attachment_number}"
            ),
            content_type="attachment",
            text=filename or "",
        )

        attachment_item["attachment_number"] = (
            attachment_number
        )

        attachment_item["filename"] = filename

        evidence.append(attachment_item)

    # Close if supported
    try:
        message.close()
    except Exception:
        pass

    return {
        "document": _relative_source(path),
        "document_type": "msg",
        "metadata": metadata,
        "attachment_count": len(attachments),
        "attachments": attachments,
        "evidence": evidence,
    }


# ============================================================
# GENERIC DOCUMENT PARSER
# ============================================================

def parse_document(source_path: str | Path) -> Dict[str, Any]:
    """
    Automatically select parser based on extension.
    """

    path = Path(source_path)

    if not path.exists():
        raise FileNotFoundError(
            f"Document does not exist: {path}"
        )

    extension = path.suffix.lower()

    if extension == ".pdf":
        return parse_pdf(path)

    if extension == ".pptx":
        return parse_pptx(path)

    if extension == ".docx":
        return parse_docx(path)

    if extension == ".msg":
        return parse_msg(path)

    raise ValueError(
        f"Unsupported document type: {extension}"
    )


# ============================================================
# DOCUMENT SEARCH
# ============================================================

def search_document(
    source_path: str | Path,
    query: str,
    content_types: Optional[List[str]] = None,
    max_results: int = 20,
) -> Dict[str, Any]:
    """
    Search within a single document.

    This is deterministic lexical retrieval.

    It does NOT use:
      - embeddings
      - vector DB
      - LLM
      - semantic similarity

    That is intentional for Step 2B.
    """

    parsed = parse_document(source_path)

    query_terms = [
        term.lower()
        for term in re.findall(
            r"\b[\w.%+-]+\b",
            query,
        )
        if term.strip()
    ]

    results = []

    for item in parsed["evidence"]:

        if (
            content_types
            and item["content_type"]
            not in content_types
        ):
            continue

        text = item.get("text", "")

        if not text:
            continue

        text_lower = text.lower()

        matched_terms = [
            term
            for term in query_terms
            if term in text_lower
        ]

        if not matched_terms:
            continue

        # Simple deterministic relevance score.
        score = len(matched_terms) / max(
            len(query_terms),
            1,
        )

        result = dict(item)

        result["matched_terms"] = matched_terms
        result["match_score"] = round(score, 4)

        results.append(result)

    results.sort(
        key=lambda x: (
            x["match_score"],
            len(x.get("matched_terms", [])),
        ),
        reverse=True,
    )

    return {
        "query": query,
        "document": parsed["document"],
        "result_count": min(
            len(results),
            max_results,
        ),
        "results": results[:max_results],
    }


# ============================================================
# READ SPECIFIC LOCATION
# ============================================================

def read_document_location(
    source_path: str | Path,
    locator: str,
) -> Dict[str, Any]:
    """
    Retrieve an exact evidence location.

    Example:

        read_document_location(
            file,
            "page=4"
        )

    or

        read_document_location(
            file,
            "slide=3"
        )
    """

    parsed = parse_document(source_path)

    exact = []

    for item in parsed["evidence"]:

        if item.get("locator") == locator:
            exact.append(item)

    return {
        "document": parsed["document"],
        "locator": locator,
        "result_count": len(exact),
        "results": exact,
    }


# ============================================================
# DOCUMENT SUMMARY
# ============================================================

def document_structure(
    source_path: str | Path,
) -> Dict[str, Any]:
    """
    Return structural information without returning all text.
    Useful for an agent deciding where to retrieve next.
    """

    parsed = parse_document(source_path)

    structure = {
        "document": parsed["document"],
        "document_type": parsed["document_type"],
        "evidence_items": len(parsed["evidence"]),
    }

    if parsed["document_type"] == "pdf":
        structure["page_count"] = parsed["page_count"]

    elif parsed["document_type"] == "pptx":
        structure["slide_count"] = parsed["slide_count"]

    elif parsed["document_type"] == "docx":
        structure["paragraph_count"] = parsed[
            "paragraph_count"
        ]
        structure["table_count"] = parsed[
            "table_count"
        ]
        structure["image_count"] = parsed[
            "image_count"
        ]

    elif parsed["document_type"] == "msg":
        structure["subject"] = parsed[
            "metadata"
        ]["subject"]

        structure["attachment_count"] = parsed[
            "attachment_count"
        ]

    return structure


# ============================================================
# JSON SERIALIZATION HELPER
# ============================================================

def save_parsed_document(
    source_path: str | Path,
    output_path: str | Path,
) -> None:
    """
    Persist parsed representation for debugging/testing.
    """

    parsed = parse_document(source_path)

    output = Path(output_path)
    output.parent.mkdir(
        parents=True,
        exist_ok=True,
    )

    with open(
        output,
        "w",
        encoding="utf-8",
    ) as f:

        json.dump(
            parsed,
            f,
            indent=2,
            ensure_ascii=False,
        )